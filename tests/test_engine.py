"""Tests for the DKCISSP Privacy Posture Assessment engine.

Unit tests (vendor detection, clause mapping) run fully offline. The
integration test hits example.com (IANA's reserved documentation/testing
domain — safe to assess, stable, and explicitly meant for this) and is
skipped gracefully if there is no network access in the test environment."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest

from privacy_posture import PrivacyPostureAssessor
from privacy_posture import vendors as V
from privacy_posture.clauses import enrich
from privacy_posture.engine import _is_safe_target, _is_blocked_ip
import ipaddress


def test_vendor_detection_google_analytics():
    html = '<script src="https://www.google-analytics.com/analytics.js"></script>'
    hits = V.detect(html)
    names = [h["name"] for h in hits]
    assert "Google Analytics" in names
    ga = next(h for h in hits if h["name"] == "Google Analytics")
    assert ga["trap_trace"] is True
    assert "GDPR" in ga["laws"]


def test_vendor_detection_cmp_not_a_tracker():
    html = '<script src="https://cdn.cookiebot.com/..."></script>'
    hits = V.detect(html)
    cmp_hits = [h for h in hits if h["is_cmp"]]
    assert cmp_hits, "CMP should be detected"
    assert cmp_hits[0]["cmp_label"] == "Cookiebot"


def test_vendor_detection_no_false_positive():
    html = "<p>Just a plain page with no third-party scripts.</p>"
    hits = V.detect(html)
    assert hits == []


def test_clause_enrich_pre_consent_tracking():
    meta = enrich("pre_consent_tracking")
    assert "GDPR" in meta["clauses"]
    assert "CCPA/CPRA" in meta["clauses"]
    assert meta["steps"], "should have remediation steps"


def test_clause_enrich_unknown_type_returns_safe_defaults():
    meta = enrich("not_a_real_finding_type")
    assert meta["clauses"] == {}
    assert meta["steps"] == []


def test_assessor_constructs_without_chromium():
    # In a CI sandbox with no Chrome/Chromium installed, this must not raise
    # and must gracefully fall back to the static scan mode.
    a = PrivacyPostureAssessor(cfg={"headless": True})
    assert isinstance(a.use_headless, bool)


def test_analyze_rejects_empty_input():
    a = PrivacyPostureAssessor()
    result = a.analyze("")
    assert "error" in result


def test_ssrf_blocks_metadata_ip_literal():
    assert _is_safe_target("169.254.169.254") is False


def test_ssrf_blocks_link_local_range():
    assert _is_safe_target("169.254.1.1") is False


def test_ssrf_blocks_userinfo_parser_differential():
    # requests/urllib3 would treat this as a request to 169.254.169.254 with
    # 'evil.com' as discarded userinfo -- our hostname-character check must
    # reject it outright rather than naively resolving 'evil.com'.
    assert _is_safe_target("evil.com@169.254.169.254") is False


def test_ssrf_blocks_unresolvable_host_fail_closed():
    assert _is_safe_target("this-domain-should-never-exist-dkcissp-test.invalid") is False


def test_ssrf_blocks_ipv4_mapped_ipv6():
    ip = ipaddress.ip_address("::ffff:169.254.169.254")
    assert _is_blocked_ip(ip) is True


def test_ssrf_blocks_nat64_embedded_metadata_ip():
    ip = ipaddress.ip_address("64:ff9b::169.254.169.254")
    assert _is_blocked_ip(ip) is True


def test_ssrf_allows_ordinary_public_ip():
    ip = ipaddress.ip_address("93.184.216.34")  # example.com-era public IP shape
    assert _is_blocked_ip(ip) is False


def test_ssrf_blocks_6to4_embedded_metadata_ip():
    # 2002:a9fe:a9fe:: is the 6to4 encoding of 169.254.169.254
    ip = ipaddress.ip_address("2002:a9fe:a9fe::")
    assert _is_blocked_ip(ip) is True


def test_ssrf_blocks_carrier_grade_nat_range():
    # 100.100.100.200 (Alibaba Cloud metadata) sits in the RFC 6598 shared
    # address space (100.64.0.0/10), which is_global correctly excludes.
    assert _is_blocked_ip(ipaddress.ip_address("100.100.100.200")) is True


def test_ssrf_blocks_loopback_and_reserved_generically():
    # The is_global-allowlist model should catch categories we never
    # enumerated by hand, not just the ones explicitly listed.
    assert _is_blocked_ip(ipaddress.ip_address("127.0.0.1")) is True
    assert _is_blocked_ip(ipaddress.ip_address("10.0.0.5")) is True
    assert _is_blocked_ip(ipaddress.ip_address("::1")) is True


def test_resolve_pinned_ip_fails_closed_on_blocked_target():
    from privacy_posture.engine import _resolve_pinned_ip
    assert _resolve_pinned_ip("169.254.169.254") is None
    assert _resolve_pinned_ip("this-domain-should-never-exist-dkcissp-test.invalid") is None
    assert _resolve_pinned_ip("evil.com@169.254.169.254") is None


def test_resolve_pinned_ip_returns_ipv4_for_safe_host():
    from privacy_posture.engine import _resolve_pinned_ip
    import socket as _socket
    try:
        _socket.getaddrinfo("example.com", None)
    except Exception:
        pytest.skip("no network access in this environment")
    pinned = _resolve_pinned_ip("example.com")
    assert pinned == "" or ipaddress.ip_address(pinned).version == 4


def test_ssrf_allows_plain_hostname_shape():
    # Shape-only check (no network call needed to prove the regex accepts a
    # normal domain) -- actual resolution is exercised by the network test.
    from privacy_posture.engine import _VALID_HOSTNAME
    assert _VALID_HOSTNAME.match("example.com")
    assert _VALID_HOSTNAME.match("sub.example.co.uk")
    assert not _VALID_HOSTNAME.match("evil.com@169.254.169.254")
    assert not _VALID_HOSTNAME.match("example.com:8080")


@pytest.mark.network
def test_analyze_example_com():
    a = PrivacyPostureAssessor(cfg={"headless": False})
    result = a.analyze("example.com")
    if result.get("error") and not result.get("reachable", True):
        pytest.skip("no network access in this environment")
    assert result["domain"] == "example.com"
    assert "score" in result
    assert "findings" in result
    assert "site" in result


def _load_report(xlsx_bytes):
    import io
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(xlsx_bytes))


def _rows(ws, min_row, max_col):
    return [[ws.cell(row=r, column=c).value for c in range(1, max_col + 1)]
            for r in range(min_row, ws.max_row + 1)]


def test_xlsx_report_with_findings():
    from privacy_posture.export import findings_to_xlsx
    result = {
        "domain": "example.com", "risk": "High", "score": 72, "scan_mode": "static",
        "site": {"https": True, "hsts": False, "banner": False, "cmp": ""},
        "kpi": {"trackers": 1, "preconsent": 0, "forms": 0, "criticals": 0},
        "laws": ["GDPR", "CCPA/CPRA"], "law_scores": {"GDPR": 80, "CCPA/CPRA": 90},
        "sev_counts": {"Critical": 0, "High": 1, "Medium": 0, "Low": 0},
        "third_parties": [{"name": "Google Analytics", "category": "Analytics",
                           "trap_trace": True, "wiretap": False}],
        "findings": [{"sev": "high", "title": "No privacy policy found/linked",
                     "detail": "detail text", "evidence": "ev text",
                     "laws": ["GDPR", "CCPA/CPRA"], "steps": ["Do X"]}],
    }
    wb = _load_report(findings_to_xlsx(result))
    assert wb.sheetnames == ["Dashboard", "Remediation Tracker", "Third Parties"]

    dash = wb["Dashboard"]
    assert "DKCISSP Privacy Posture Assessment" in dash["B2"].value
    assert "example.com" in dash["B4"].value
    assert any(row[1] == "GDPR" and row[2] == 80 for row in _rows(dash, 12, 3))

    rt = wb["Remediation Tracker"]
    rt_rows = _rows(rt, 4, 12)
    task_rows = [r for r in rt_rows if r[2] == "No privacy policy found/linked"]
    # exactly one task row for the one finding's one remediation step -- no
    # truncation, no extra rows
    assert len(task_rows) == 1
    assert task_rows[0][1] == "High"
    assert task_rows[0][3] == "Do X"
    assert task_rows[0][4] == "GDPR; CCPA/CPRA"
    assert task_rows[0][6] == "Not Started"

    tp = wb["Third Parties"]
    assert any(row[0] == "Google Analytics" for row in _rows(tp, 3, 5))


def test_xlsx_report_splits_multiple_steps_into_one_row_each():
    """A finding with several remediation steps must become one Remediation
    Tracker row per step, not one row with everything crammed together, so
    each step can be tracked (Owner/Status/Due Date) independently."""
    from privacy_posture.export import findings_to_xlsx
    result = {
        "domain": "example.com", "risk": "Critical", "score": 10, "scan_mode": "dynamic",
        "site": {"https": True, "hsts": True, "banner": True, "cmp": "OneTrust"},
        "kpi": {"trackers": 2, "preconsent": 1, "forms": 0, "criticals": 1},
        "laws": ["GDPR"], "law_scores": {"GDPR": 40},
        "sev_counts": {"Critical": 1, "High": 1, "Medium": 0, "Low": 0},
        "third_parties": [],
        "findings": [
            {"sev": "critical", "title": "Finding A",
             "detail": "line1\nline2", "evidence": "beacon1\nbeacon2\nbeacon3",
             "laws": ["GDPR"], "steps": ["step one", "step two"]},
            {"sev": "high", "title": "Finding B", "detail": "d", "evidence": "e",
             "laws": ["GDPR"], "steps": ["only step"]},
        ],
    }
    wb = _load_report(findings_to_xlsx(result))
    rt = wb["Remediation Tracker"]
    rows = _rows(rt, 4, 12)
    finding_a_rows = [r for r in rows if r[2] == "Finding A"]
    finding_b_rows = [r for r in rows if r[2] == "Finding B"]
    assert len(finding_a_rows) == 2
    assert len(finding_b_rows) == 1
    assert {r[3] for r in finding_a_rows} == {"step one", "step two"}
    # High-severity default priority, Critical gets the top default priority
    assert all(r[7] == "P1" for r in finding_a_rows)
    assert finding_b_rows[0][7] == "P2"


def test_xlsx_report_with_no_findings():
    from privacy_posture.export import findings_to_xlsx
    result = {"domain": "example.com", "risk": "Low", "score": 100,
             "scan_mode": "static", "findings": []}
    wb = _load_report(findings_to_xlsx(result))
    rt = wb["Remediation Tracker"]
    rows = _rows(rt, 4, 12)
    assert any(r[2] == "No privacy issues detected" for r in rows)


def test_cli_entrypoint_exists():
    from privacy_posture.cli import main
    assert callable(main)
